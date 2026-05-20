import openpyxl

# File 1: Test Status
wb = openpyxl.load_workbook(r'D:\zPython\MakerAi\Docs\Pruebas\MakerAI34_Test_Status_Feb_25_2026.xlsx', data_only=True)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f'\n=== {sn} ===')
    for row in ws.iter_rows(max_row=5, values_only=True):
        vals = [str(v) for v in row if v is not None]
        if vals:
            print(' | '.join(vals[:8])[:200])

# File 2: Test List
print('\n\n=== Test List.xlsx ===')
wb2 = openpyxl.load_workbook(r'D:\zPython\MakerAi\Docs\Version 3\EN\Test List.xlsx', data_only=True)
for sn in wb2.sheetnames:
    ws = wb2[sn]
    print(f'\n--- {sn} ---')
    for row in ws.iter_rows(max_row=5, values_only=True):
        vals = [str(v) for v in row if v is not None]
        if vals:
            print(' | '.join(vals[:8])[:200])

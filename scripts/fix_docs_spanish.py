#!/usr/bin/env python3
"""Fix remaining Spanish fragments in EN .docx and .xlsx files."""
import docx, openpyxl, os, re, shutil

ROOT = r'D:\zPython\MakerAi'

# =============================================
# 1. Fix Installation_EN.docx - WORST CASE
# =============================================
INSTALL_PATH = os.path.join(ROOT, r'Docs\Version 3\EN\Manual-Installation\Manual-Installation_EN.docx')

# Specific Spanish->English replacements
INSTALL_FIXES = {
    'Requisitos previos': 'Prerequisites',
    'Tener instalado Delphi': 'Have Delphi installed',
    'en tu sistema': 'on your system',
    'Una conexión a internet para clonar el repositorio': 'An internet connection to clone the repository',
    'Opcional: Tener instalado Git': 'Optional: Have Git installed',
    'para clonar el repositorio': 'to clone the repository',
    'Opcional': 'Optional',
}

def fix_docx(filepath, replacements, label):
    """Apply text replacements to a .docx file."""
    if not os.path.exists(filepath):
        print(f'  [SKIP] {label}: file not found')
        return 0
    
    # Backup
    backup = filepath + '.bak'
    shutil.copy2(filepath, backup)
    
    doc = docx.Document(filepath)
    changes = 0
    
    for para in doc.paragraphs:
        orig = para.text
        text = orig
        for old, new in replacements.items():
            if old in text:
                text = text.replace(old, new)
        if text != orig:
            # Clear and rewrite the paragraph
            para.clear()
            run = para.add_run(text)
            changes += 1
    
    # Also fix tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    orig = para.text
                    text = orig
                    for old, new in replacements.items():
                        if old in text:
                            text = text.replace(old, new)
                    if text != orig:
                        para.clear()
                        run = para.add_run(text)
                        changes += 1
    
    if changes > 0:
        doc.save(filepath)
        os.remove(backup)
        print(f'  [FIXED] {label}: {changes} paragraphs')
    else:
        os.remove(backup)
        print(f'  [OK] {label}: no changes needed')
    return changes

# Fix Installation
fix_docx(INSTALL_PATH, INSTALL_FIXES, 'Installation_EN.docx')

# =============================================
# 2. Fix other .docx files with minor fragments
# =============================================

DOCX_FIXES = {
    # Agents_EN.docx
    r'Docs\Version 3\EN\Agents\uMakerAi-Agents\uMakerAi-Agents_EN.docx': {
        'Ahora necesitamos decirle a nuestro': 'Now we need to tell our',
        'Ahora puedes reutilizar': 'Now you can reuse',
    },
    # Agents_AutoAgents_EN.docx
    r'Docs\Version 3\EN\Agents\uMakerAi-AutoAgents-D02-SingleTool\uMakerAi-AutoAgents-D02-SingleTool_EN.docx': {
        'categoria opcional': 'optional category',
        'GetMessages() incluye ahora:': 'GetMessages() now includes:',
    },
    # Annexes_EN.docx
    r'Docs\Version 3\EN\Agents\Annexes-ConditionsLanguage_EN.docx': {
        'Extensiones futuras (opcional)': 'Future extensions (optional)',
    },
    # MCP.Server.EN.docx
    r'Docs\Version 3\EN\MCPServer\uMakerAi-MCP.Server.EN.docx': {
        'Configuración para exponer un sistema de archivos local vía MCP:': 'Configuration for exposing a local file system via MCP:',
    },
    # RAG_EN.docx
    r'Docs\Version 3\EN\RAG\uMakerAi-RAG_EN.docx': {
        'Ahora, ¡a probar!': "Now, let's test!",
    },
    # ChatConnection_EN.docx
    r'Docs\Version 3\EN\uMakerAi-ChatConnection\uMakerAi-ChatConnection_EN.docx': {
        'Ahora TAiChatConnection': 'Now TAiChatConnection',
    },
    # Prompts_EN.docx
    r'Docs\Version 3\EN\uMakerAi.Prompts\uMakerAi.Prompts_EN.docx': {
        '// Ahora PromptFinal': '// Now PromptFinal',
    },
}

for relpath, replacements in DOCX_FIXES.items():
    filepath = os.path.join(ROOT, relpath)
    label = os.path.basename(filepath)
    fix_docx(filepath, replacements, label)

# =============================================
# 3. Fix .xlsx files - Spanish headers
# =============================================

XLSX_FILES = {
    # Test Status
    r'Docs\Pruebas\MakerAI34_Test_Status_Feb_25_2026.xlsx': {
        'Categoría': 'Category',
        'Característica': 'Feature',
        'Proveedor': 'Provider',
        'Modelo Utilizado': 'Model Used',
        'Estado Pruebas': 'Test Status',
        'Modelos por Característica': 'Models by Feature',
    },
    # Test List
    r'Docs\Version 3\EN\Test List.xlsx': {
        'Sincrono': 'Synchronous',
        'Asincrono': 'Asynchronous',
        'Hoja1': 'Sheet1',
        'Hoja2': 'Sheet2',
    },
}

for relpath, replacements in XLSX_FILES.items():
    filepath = os.path.join(ROOT, relpath)
    if not os.path.exists(filepath):
        print(f'  [SKIP] {os.path.basename(filepath)}: not found')
        continue
    
    backup = filepath + '.bak'
    shutil.copy2(filepath, backup)
    
    wb = openpyxl.load_workbook(filepath)
    changes = 0
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # Fix sheet name itself
        if sheet_name in replacements:
            new_name = replacements[sheet_name]
            ws.title = new_name
            changes += 1
            print(f'  Sheet: "{sheet_name}" -> "{new_name}"')
        
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    orig = cell.value
                    text = orig
                    for old, new in replacements.items():
                        if old in text:
                            text = text.replace(old, new)
                    if text != orig:
                        cell.value = text
                        changes += 1
    
    if changes > 0:
        wb.save(filepath)
        os.remove(backup)
        print(f'  [FIXED] {os.path.basename(filepath)}: {changes} cell changes')
    else:
        os.remove(backup)
        print(f'  [OK] {os.path.basename(filepath)}: no changes')

print('\nDone fixing docs. Running verification...')

# =============================================
# 4. Verify fixes
# =============================================
print('\n=== Verification ===')

# Re-run the check script
import subprocess
result = subprocess.run(['python', os.path.join(ROOT, r'scripts\check_en_docs.py')], 
                       capture_output=True, text=True, cwd=ROOT)
lines = result.stdout.split('\n')
leak_count = sum(1 for l in lines if '[ES-LEAK]' in l)
print(f'Remaining leaks in EN docs: {leak_count}')
if leak_count > 0:
    for l in lines:
        if '[ES-LEAK]' in l or 'Leaking' in l:
            print(f'  {l.strip()}')

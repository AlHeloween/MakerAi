import docx
doc = docx.Document(r'D:\zPython\MakerAi\Docs\Version 3\EN\uMakerAi-ChatConnection\uMakerAi-ChatConnection_EN.docx')
for p in doc.paragraphs:
    t = p.text
    if 'Ahora' in t:
        print(repr(t[:250]))
        break
else:
    print('Ahora not found in paragraphs')
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p2 in cell.paragraphs:
                    t = p2.text
                    if 'Ahora' in t:
                        print(f'FOUND in table cell: {repr(t[:250])}')
                        break

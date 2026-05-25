import openpyxl
wb = openpyxl.load_workbook('D:/zPython/MakerAi/Docs/Pruebas/MakerAI34_Test_Status_Feb_25_2026.xlsx', data_only=True)

print('=== Test Status (corrected cells) ===')
ws = wb['Test Status']
check_cells = ['H11','F14','L14','G12','G15','H16','H17']
for row in ws.iter_rows(min_row=5, max_row=17, values_only=False):
    cells = []
    for c in row:
        v = c.value or ''
        marker = ' *' if c.coordinate in check_cells else ''
        cells.append(f'{str(v):16s}{marker}')
    if any(str(c.value or '').strip() for c in row) and '*' in ''.join(cells):
        print(' | '.join(cells))

print()
print('=== Models by Feature (new entries row 19+) ===')
ws2 = wb['Models by Feature']
for row in ws2.iter_rows(min_row=19, values_only=False):
    cells = []
    for c in row:
        v = c.value or ''
        cells.append(f'{c.coordinate}={repr(v)}')
    if cells and any(c.value for c in row):
        print(' | '.join(cells))
